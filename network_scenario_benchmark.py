#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Network x Scenario Benchmark: 4 Models x 10 Networks x 3 Demand Scenarios
==========================================================================

Runs each of 4 inventory policies on each of 10 supply chain topologies under
3 demand scenarios and writes a single Excel report comparing average total
cost and fill rate, with gap-% versus the proposed GNN-HAPPO.

Models       : (s,S) BaseStock, MAPPO, HAPPO, GNN-HAPPO (proposed)
Networks     : 1x3, 1x7, 1x10, 2x15, 2x20, 2x30, 2x40, 4x15, 4x30, 4x40
Scenarios    : S1-Balanced, S2-High Demand, S3-Extreme Stress
Episodes     : 5 per (model, network, scenario)

Per-network evaluator classes from `test_trained_model_*_<net>.py` and
`Test_baseline_basestock_<net>.py` are imported dynamically; demand scenarios
are injected into the env after construction (mean / std per SKU).

Output:
  Network_Scenario_Benchmark.xlsx with sheets:
    - Overview        : matrices of mean total cost & fill rate per scenario
    - S1_Balanced     : per-network 4-model comparison + gap-% vs GNN-HAPPO
    - S2_HighDemand   : same
    - S3_Extreme      : same
    - Gap_Summary     : average gap-% per baseline across networks/scenarios
    - Raw_Data        : flat table (model, network, scenario, episode, metrics)
    - Methodology     : formulas + notes

Usage
-----
    python network_scenario_benchmark.py
    python network_scenario_benchmark.py --num_episodes 5 --episode_length 365
    python network_scenario_benchmark.py --models gnn happo --networks 1x3 1x7
"""

from __future__ import annotations

import argparse
import gc
import importlib
import sys
import traceback
from argparse import Namespace
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
import torch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# ---------------------------------------------------------------------------
# Project root on sys.path so dynamic imports of test_trained_model_* work.
# ---------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


# ---------------------------------------------------------------------------
# Static configuration
# ---------------------------------------------------------------------------

NETWORKS: list[str] = [
    "1x3", "1x7", "1x10",
    "2x15", "2x20", "2x30", "2x40",
    "4x15", "4x30", "4x40",
]

MODEL_FAMILIES = ["basestock", "mappo", "happo", "gnn"]
MODEL_LABELS = {
    "basestock": "(s,S) BaseStock",
    "mappo":     "MAPPO",
    "happo":     "HAPPO",
    "gnn":       "GNN-HAPPO",
}
BASELINES = ["basestock", "mappo", "happo"]
PROPOSED = "gnn"

DEMAND_SCENARIOS: dict[str, dict] = {
    "Scenario_1_Balanced": {
        "SKU_0": {"mean": 1.41,  "std": 1.99},
        "SKU_1": {"mean": 1.06,  "std": 1.28},
        "SKU_2": {"mean": 0.77,  "std": 1.06},
        "short": "S1-Balanced",
        "label": "S1: Balanced (Low Stress)",
        "sheet": "S1_Balanced",
    },
    "Scenario_2_HighDemand": {
        "SKU_0": {"mean": 1.81,  "std": 1.99},
        "SKU_1": {"mean": 1.378, "std": 1.28},
        "SKU_2": {"mean": 1.001, "std": 1.06},
        "short": "S2-High",
        "label": "S2: High Demand (Mixed Stress)",
        "sheet": "S2_HighDemand",
    },
    "Scenario_3_HighVariance": {
        "SKU_0": {"mean": 1.41,  "std": 2.0},
        "SKU_1": {"mean": 1.06,  "std": 1.4},
        "SKU_2": {"mean": 0.77,  "std": 1.2},
        "short": "S3-Extreme",
        "label": "S3: Extreme Stress (High Volatility)",
        "sheet": "S3_Extreme",
    },
}

# (s,S) policy reorder points used across all networks (same defaults as
# Test_baseline_basestock_*.py).
SS_PARAMS = dict(s_dc=100.0, S_dc=170.0, s_retailer=3.0, S_retailer=10.0)


# ---------------------------------------------------------------------------
# Per-(model, network) class & model-dir resolution
# ---------------------------------------------------------------------------

def evaluator_class_for(model: str, net: str):
    """Dynamically import the right Evaluator class for (model, net).

    2x15 is a special case: its file holds the *base* class (no network
    suffix in the class name).
    """
    if model == "basestock":
        if net == "2x15":
            module = importlib.import_module("Test_baseline_basestock_2x15")
            return module.BaseStockEvaluator
        module = importlib.import_module(f"Test_baseline_basestock_{net}")
        return getattr(module, f"BaseStockEvaluator{net}")
    if model == "gnn":
        if net == "2x15":
            module = importlib.import_module("test_trained_model_gnn_2x15")
            return module.GNNModelEvaluator
        module = importlib.import_module(f"test_trained_model_gnn_{net}")
        return getattr(module, f"GNNModelEvaluator{net}")
    if model == "happo":
        if net == "2x15":
            module = importlib.import_module("test_trained_model_happo_2x15")
            return module.ModelEvaluator
        module = importlib.import_module(f"test_trained_model_happo_{net}")
        return getattr(module, f"HAPPOEvaluator{net}")
    if model == "mappo":
        if net == "2x15":
            module = importlib.import_module("test_trained_model_mappo_2x15")
            return module.MAPPOModelEvaluator
        module = importlib.import_module(f"test_trained_model_mappo_{net}")
        return getattr(module, f"MAPPOEvaluator{net}")
    raise ValueError(f"Unknown model family {model!r}")


def model_dir_for(model: str, net: str) -> Path | None:
    """Standard model directory pattern. Returns None for (s,S) BaseStock."""
    if model == "basestock":
        return None
    if model == "gnn":
        # 2x15 was trained by train_multi_dc_gnn.py -> experiment "gnn_happo_full"
        if net == "2x15":
            return ROOT / "results" / "gnn_happo_full" / "run_seed_1" / "models"
        return ROOT / "results" / f"gnn_happo_{net}" / "run_seed_1" / "models"
    if model == "happo":
        # 2x15 was trained by train_multi_dc_baseline.py -> experiment "full_training"
        if net == "2x15":
            return ROOT / "results" / "full_training" / "run_seed_1" / "models"
        return ROOT / "results" / f"happo_{net}" / "run_seed_1" / "models"
    if model == "mappo":
        return ROOT / "results" / f"mappo_{net}" / "run_seed_1" / "models"
    raise ValueError(f"Unknown model family {model!r}")


def config_path_for(net: str) -> str:
    return f"configs/multi_dc_{net}_config.yaml"


# ---------------------------------------------------------------------------
# Scenario injection
# ---------------------------------------------------------------------------

def _apply_scenario_to_env(base_env, scenario: dict) -> None:
    """Inject per-SKU mean/std into a MultiDCInventoryEnv instance."""
    n_skus = getattr(base_env, "n_skus", 3)
    new_mean = np.array(base_env.demand_mean, dtype=float).copy()
    new_std  = np.array(base_env.demand_std,  dtype=float).copy()
    for sku_idx in range(n_skus):
        key = f"SKU_{sku_idx}"
        if key in scenario:
            new_mean[sku_idx] = scenario[key]["mean"]
            new_std[sku_idx]  = scenario[key]["std"]
    base_env.demand_mean = new_mean
    base_env.demand_std  = new_std


def apply_scenario(evaluator, scenario: dict) -> None:
    """Apply demand scenario to either a raw env or a vectorised env."""
    env = getattr(evaluator, "env", None)
    if env is None:
        return
    env_list = getattr(env, "env_list", None) or getattr(env, "envs", None)
    if env_list:
        for sub in env_list:
            _apply_scenario_to_env(sub, scenario)
    else:
        _apply_scenario_to_env(env, scenario)


# ---------------------------------------------------------------------------
# Args construction
# ---------------------------------------------------------------------------

def build_args(model: str, net: str, episode_length: int, num_episodes: int,
               seed: int, cuda: bool, model_dir: Path | None) -> Namespace:
    """Build an argparse.Namespace with everything the various Evaluator
    classes might read off `args`."""
    args = Namespace()
    # Shared
    args.episode_length = episode_length
    args.num_episodes   = num_episodes
    args.seed           = seed
    args.cuda           = cuda
    args.save_dir       = str(ROOT / "evaluation_results" / "network_scenario_benchmark_tmp")
    args.experiment_name = f"{model}_{net}_eval"

    # BaseStock (s,S) parameters & config path
    args.config_path = config_path_for(net)
    for k, v in SS_PARAMS.items():
        setattr(args, k, v)

    # RL-evaluator fields
    if model_dir is not None:
        args.model_dir = str(model_dir)
    args.algorithm_name = {
        "happo": "happo",
        "mappo": "mappo",
        "gnn":   "gnn_happo",
        "basestock": "ss_heuristic",
    }[model]

    # 2x15 specifically needs num_agents because its base class reads it
    # before _create_env populates it.
    net_to_agents = {
        "1x3": 4, "1x7": 8, "1x10": 11,
        "2x15": 17, "2x20": 22, "2x30": 32, "2x40": 42,
        "4x15": 19, "4x30": 34, "4x40": 44,
    }
    args.num_agents = net_to_agents[net]

    # GNN architecture defaults (consumed only by GNN evaluators).
    args.gnn_type = "GAT"
    args.gnn_hidden_dim = 128
    args.gnn_num_layers = 2
    args.num_attention_heads = 4
    args.gnn_dropout = 0.1
    args.use_residual = True
    args.critic_pooling = "mean"

    return args


# ---------------------------------------------------------------------------
# Metric extraction from evaluator.episode_metrics
# ---------------------------------------------------------------------------

def compute_episode_metrics(evaluator) -> list[dict]:
    """For each completed episode, compute Total_Cost, Fill_Rate, Lost_Sales,
    Avg_Inventory, and cost decomposition — same formula used by the
    per-network test scripts' `_save_metrics_csv`, but with the correct
    `n_dcs` for this topology (instead of hardcoded 2)."""
    out = []
    n_agents = int(getattr(evaluator, "n_agents",
                           getattr(evaluator.args, "num_agents", 0)))
    n_dcs = int(getattr(evaluator, "n_dcs", 0))
    if n_dcs == 0:
        # Recover from the env if attribute wasn't set.
        env = evaluator.env
        sub = (getattr(env, "env_list", None) or getattr(env, "envs", None) or [env])[0]
        n_dcs = int(getattr(sub, "n_dcs", 0))
    for m in evaluator.episode_metrics:
        total_holding  = float(np.sum(m["holding_costs"]))
        total_backlog  = float(np.sum(m["backlog_costs"]))
        total_ordering = float(np.sum(m["ordering_costs"]))
        total_cost     = total_holding + total_backlog + total_ordering

        placed     = sum(m["_orders_placed"][aid]     for aid in range(n_dcs, n_agents))
        from_stock = sum(m["_orders_from_stock"][aid] for aid in range(n_dcs, n_agents))
        lost_sales = max(placed - from_stock, 0)
        fill_rate = float(np.mean(m["service_level"]))
        avg_inventory = float(np.mean(m["avg_inventory"]))

        out.append({
            "Total_Cost":          total_cost,
            "Fill_Rate":           fill_rate,
            "Lost_Sales":          lost_sales,
            "Avg_Inventory":       avg_inventory,
            "Total_Holding_Cost":  total_holding,
            "Total_Backlog_Cost":  total_backlog,
            "Total_Ordering_Cost": total_ordering,
        })
    return out


# ---------------------------------------------------------------------------
# Evaluator orchestration
# ---------------------------------------------------------------------------

def run_one(model: str, net: str, scenario_key: str, scenario: dict,
            episode_length: int, num_episodes: int, seed: int,
            cuda: bool) -> dict:
    """Run one (model, network, scenario) combination and return a dict
    with status + per-episode rows. Skips RL models whose checkpoint is
    missing (status='MISSING')."""

    model_dir = model_dir_for(model, net)
    if model != "basestock" and (model_dir is None or not model_dir.exists()):
        return {
            "status": "MISSING",
            "reason": f"Model directory not found: {model_dir}",
            "episodes": [],
        }

    EvaluatorCls = evaluator_class_for(model, net)
    args = build_args(model, net, episode_length, num_episodes, seed, cuda, model_dir)

    print(f"\n--- {MODEL_LABELS[model]:<18} | net={net:<5} | {scenario['short']} ---")
    try:
        evaluator = EvaluatorCls(args)
    except Exception as exc:
        print(f"  [SKIP] Evaluator construction failed: {exc}")
        return {"status": "ERROR", "reason": str(exc), "episodes": []}

    try:
        apply_scenario(evaluator, scenario)
        evaluator.evaluate()
        episodes = compute_episode_metrics(evaluator)
    except Exception as exc:
        traceback.print_exc()
        return {"status": "ERROR", "reason": str(exc), "episodes": []}
    finally:
        # release env + policies between runs
        del evaluator
        gc.collect()
        if torch.cuda.is_available():
            torch.cuda.empty_cache()

    return {"status": "OK", "episodes": episodes}


def run_all(models: list[str], networks: list[str], scenarios: dict,
            episode_length: int, num_episodes: int, seed: int,
            cuda: bool) -> pd.DataFrame:
    """Run the full grid and return a flat DataFrame of per-episode rows.

    Columns: Model, Network, Scenario, Episode, Status,
             Total_Cost, Fill_Rate, Lost_Sales, Avg_Inventory,
             Total_Holding_Cost, Total_Backlog_Cost, Total_Ordering_Cost
    """
    rows: list[dict] = []
    total = len(models) * len(networks) * len(scenarios)
    done = 0
    for model in models:
        for net in networks:
            for sc_key, sc_val in scenarios.items():
                done += 1
                print(f"\n[{done}/{total}] {MODEL_LABELS[model]} on {net} ({sc_val['short']})")
                result = run_one(
                    model, net, sc_key, sc_val,
                    episode_length, num_episodes, seed, cuda,
                )
                if result["status"] != "OK":
                    rows.append({
                        "Model": MODEL_LABELS[model],
                        "Network": net,
                        "Scenario": sc_val["short"],
                        "Episode": 0,
                        "Status": result["status"],
                        "Note": result.get("reason", ""),
                        "Total_Cost": np.nan, "Fill_Rate": np.nan,
                        "Lost_Sales": np.nan, "Avg_Inventory": np.nan,
                        "Total_Holding_Cost": np.nan,
                        "Total_Backlog_Cost": np.nan,
                        "Total_Ordering_Cost": np.nan,
                    })
                    continue
                for i, ep in enumerate(result["episodes"], start=1):
                    rows.append({
                        "Model": MODEL_LABELS[model],
                        "Network": net,
                        "Scenario": sc_val["short"],
                        "Episode": i,
                        "Status": "OK",
                        "Note": "",
                        **ep,
                    })
    return pd.DataFrame(rows)


# ===========================================================================
# Excel report builder
# ===========================================================================

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
GNN_FILL = PatternFill("solid", fgColor="E2EFDA")
BEST_FILL = PatternFill("solid", fgColor="FFE699")
MISS_FILL = PatternFill("solid", fgColor="F4CCCC")
POS_FONT = Font(color="006100", bold=True)
NEG_FONT = Font(color="9C0006", bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws: Worksheet, row: int, n_cols: int) -> None:
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def autosize(ws: Worksheet, widths: Iterable[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def aggregate(df: pd.DataFrame) -> pd.DataFrame:
    """Aggregate per-episode rows -> mean per (Model, Network, Scenario)."""
    if df.empty:
        return df
    ok = df[df["Status"] == "OK"].copy()
    if ok.empty:
        return ok
    grp = ok.groupby(["Model", "Network", "Scenario"], as_index=False)
    agg = grp.agg(
        Mean_Total_Cost=("Total_Cost", "mean"),
        Std_Total_Cost =("Total_Cost", "std"),
        Mean_Fill_Rate =("Fill_Rate", "mean"),
        Std_Fill_Rate  =("Fill_Rate", "std"),
        Mean_Lost_Sales=("Lost_Sales", "mean"),
        Mean_Avg_Inventory=("Avg_Inventory", "mean"),
        N_Episodes=("Episode", "count"),
    )
    return agg


def lookup(agg: pd.DataFrame, model_label: str, net: str, scen: str,
           col: str) -> float:
    sub = agg[(agg["Model"] == model_label) & (agg["Network"] == net) &
              (agg["Scenario"] == scen)]
    if sub.empty:
        return float("nan")
    return float(sub.iloc[0][col])


def build_overview(ws: Worksheet, agg: pd.DataFrame, networks: list[str],
                   scenarios: dict) -> None:
    ws.cell(row=1, column=1,
            value="Network x Scenario Benchmark -- Overview").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1,
            value=("Mean across 5 evaluation episodes per (model, network, "
                   "scenario). Lower cost / higher fill rate is better.")
            ).font = Font(italic=True, color="595959")

    model_labels = [MODEL_LABELS[m] for m in MODEL_FAMILIES]
    row_cursor = 4

    for metric_name, col_in_agg, fmt, better in [
        ("Mean Total Cost (VND thousands)", "Mean_Total_Cost", "#,##0.00", "min"),
        ("Mean Fill Rate (%)", "Mean_Fill_Rate", "0.00", "max"),
    ]:
        for sc_val in scenarios.values():
            scen = sc_val["short"]
            ws.cell(row=row_cursor, column=1,
                    value=f"{metric_name}  |  {sc_val['label']}"
                    ).font = Font(bold=True, size=12)
            headers = ["Network", *model_labels, "Best Model"]
            for j, h in enumerate(headers, start=1):
                ws.cell(row=row_cursor + 1, column=j, value=h)
            style_header(ws, row_cursor + 1, len(headers))

            for i, net in enumerate(networks):
                r = row_cursor + 2 + i
                ws.cell(row=r, column=1, value=net).alignment = CENTER
                vals: list[float] = []
                for j, m in enumerate(MODEL_FAMILIES, start=2):
                    v = lookup(agg, MODEL_LABELS[m], net, scen, col_in_agg)
                    if not np.isnan(v) and col_in_agg == "Mean_Total_Cost":
                        v = v / 1000.0
                    vals.append(v)
                    cell = ws.cell(row=r, column=j, value=(None if np.isnan(v) else v))
                    cell.number_format = fmt
                    cell.alignment = RIGHT
                    if np.isnan(v):
                        cell.fill = MISS_FILL
                # Best column
                finite = [(j, v) for j, v in enumerate(vals) if not np.isnan(v)]
                if finite:
                    if better == "min":
                        best_idx, _ = min(finite, key=lambda x: x[1])
                    else:
                        best_idx, _ = max(finite, key=lambda x: x[1])
                    ws.cell(row=r, column=best_idx + 2).fill = BEST_FILL
                    ws.cell(row=r, column=2 + len(MODEL_FAMILIES),
                            value=model_labels[best_idx]).alignment = CENTER
                # Highlight GNN column
                gnn_col = 2 + MODEL_FAMILIES.index(PROPOSED)
                ws.cell(row=r, column=gnn_col).fill = GNN_FILL
                for c in range(1, len(headers) + 1):
                    ws.cell(row=r, column=c).border = BORDER

            row_cursor += 2 + len(networks) + 2

    autosize(ws, [10, 18, 18, 18, 18, 18])
    ws.freeze_panes = "B5"


def build_scenario_sheet(ws: Worksheet, agg: pd.DataFrame, networks: list[str],
                         scen_short: str, scen_label: str) -> None:
    ws.cell(row=1, column=1,
            value=f"Scenario: {scen_label}").font = Font(bold=True, size=14)

    headers = [
        "Network",
        "(s,S) Cost (VND k)", "MAPPO Cost (VND k)", "HAPPO Cost (VND k)",
        "GNN-HAPPO Cost (VND k)",
        "(s,S) Gap %", "MAPPO Gap %", "HAPPO Gap %",
        "(s,S) Fill %", "MAPPO Fill %", "HAPPO Fill %", "GNN-HAPPO Fill %",
        "(s,S) FillGap pp", "MAPPO FillGap pp", "HAPPO FillGap pp",
    ]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=3, column=j, value=h)
    style_header(ws, 3, len(headers))

    def gap_cost(b_cost: float, g_cost: float) -> float:
        if np.isnan(b_cost) or np.isnan(g_cost) or b_cost == 0:
            return float("nan")
        return (b_cost - g_cost) / b_cost * 100.0

    def gap_fill_pp(g_fill: float, b_fill: float) -> float:
        if np.isnan(g_fill) or np.isnan(b_fill):
            return float("nan")
        return g_fill - b_fill

    for i, net in enumerate(networks):
        r = 4 + i
        ws.cell(row=r, column=1, value=net).alignment = CENTER

        # Costs
        bs_c = lookup(agg, MODEL_LABELS["basestock"], net, scen_short, "Mean_Total_Cost")
        ma_c = lookup(agg, MODEL_LABELS["mappo"],     net, scen_short, "Mean_Total_Cost")
        ha_c = lookup(agg, MODEL_LABELS["happo"],     net, scen_short, "Mean_Total_Cost")
        gn_c = lookup(agg, MODEL_LABELS["gnn"],       net, scen_short, "Mean_Total_Cost")
        for col, v in zip([2, 3, 4, 5], [bs_c, ma_c, ha_c, gn_c]):
            cell = ws.cell(row=r, column=col, value=(None if np.isnan(v) else v / 1000.0))
            cell.number_format = "#,##0.00"
            cell.alignment = RIGHT
            if np.isnan(v):
                cell.fill = MISS_FILL
        ws.cell(row=r, column=5).fill = GNN_FILL  # highlight GNN column

        # Gap % vs GNN-HAPPO  (positive = GNN cheaper)
        for col, base in zip([6, 7, 8], [bs_c, ma_c, ha_c]):
            v = gap_cost(base, gn_c)
            cell = ws.cell(row=r, column=col, value=(None if np.isnan(v) else v))
            cell.number_format = "0.00"
            cell.alignment = RIGHT
            if not np.isnan(v):
                cell.font = POS_FONT if v >= 0 else NEG_FONT
            else:
                cell.fill = MISS_FILL

        # Fill rates
        bs_f = lookup(agg, MODEL_LABELS["basestock"], net, scen_short, "Mean_Fill_Rate")
        ma_f = lookup(agg, MODEL_LABELS["mappo"],     net, scen_short, "Mean_Fill_Rate")
        ha_f = lookup(agg, MODEL_LABELS["happo"],     net, scen_short, "Mean_Fill_Rate")
        gn_f = lookup(agg, MODEL_LABELS["gnn"],       net, scen_short, "Mean_Fill_Rate")
        for col, v in zip([9, 10, 11, 12], [bs_f, ma_f, ha_f, gn_f]):
            cell = ws.cell(row=r, column=col, value=(None if np.isnan(v) else v))
            cell.number_format = "0.00"
            cell.alignment = RIGHT
            if np.isnan(v):
                cell.fill = MISS_FILL
        ws.cell(row=r, column=12).fill = GNN_FILL

        # Fill gap (pp)  (positive = GNN serves more demand)
        for col, base in zip([13, 14, 15], [bs_f, ma_f, ha_f]):
            v = gap_fill_pp(gn_f, base)
            cell = ws.cell(row=r, column=col, value=(None if np.isnan(v) else v))
            cell.number_format = "0.00"
            cell.alignment = RIGHT
            if not np.isnan(v):
                cell.font = POS_FONT if v >= 0 else NEG_FONT
            else:
                cell.fill = MISS_FILL

        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = BORDER

    # Avg-across-networks row
    rr = 4 + len(networks) + 1
    ws.cell(row=rr, column=1, value="Average").font = Font(bold=True)
    style_header(ws, rr, len(headers))
    rr += 1
    ws.cell(row=rr, column=1, value="ALL").alignment = CENTER

    def avg_lookup(col: str, model_key: str) -> float:
        vals = [lookup(agg, MODEL_LABELS[model_key], n, scen_short, col)
                for n in networks]
        vals = [v for v in vals if not np.isnan(v)]
        return float(np.mean(vals)) if vals else float("nan")

    bs_c_avg = avg_lookup("Mean_Total_Cost", "basestock")
    ma_c_avg = avg_lookup("Mean_Total_Cost", "mappo")
    ha_c_avg = avg_lookup("Mean_Total_Cost", "happo")
    gn_c_avg = avg_lookup("Mean_Total_Cost", "gnn")
    for col, v in zip([2, 3, 4, 5], [bs_c_avg, ma_c_avg, ha_c_avg, gn_c_avg]):
        cell = ws.cell(row=rr, column=col, value=(None if np.isnan(v) else v / 1000.0))
        cell.number_format = "#,##0.00"
        cell.alignment = RIGHT
    ws.cell(row=rr, column=5).fill = GNN_FILL

    # Mean of per-network relative gaps (equal weight per network)
    def per_network_mean_gap(model_key: str) -> float:
        diffs = []
        for n in networks:
            b = lookup(agg, MODEL_LABELS[model_key], n, scen_short, "Mean_Total_Cost")
            g = lookup(agg, MODEL_LABELS["gnn"],     n, scen_short, "Mean_Total_Cost")
            if np.isnan(b) or np.isnan(g) or b == 0:
                continue
            diffs.append((b - g) / b * 100.0)
        return float(np.mean(diffs)) if diffs else float("nan")

    for col, mk in zip([6, 7, 8], ["basestock", "mappo", "happo"]):
        v = per_network_mean_gap(mk)
        cell = ws.cell(row=rr, column=col, value=(None if np.isnan(v) else v))
        cell.number_format = "0.00"
        cell.alignment = RIGHT
        if not np.isnan(v):
            cell.font = POS_FONT if v >= 0 else NEG_FONT

    bs_f_avg = avg_lookup("Mean_Fill_Rate", "basestock")
    ma_f_avg = avg_lookup("Mean_Fill_Rate", "mappo")
    ha_f_avg = avg_lookup("Mean_Fill_Rate", "happo")
    gn_f_avg = avg_lookup("Mean_Fill_Rate", "gnn")
    for col, v in zip([9, 10, 11, 12], [bs_f_avg, ma_f_avg, ha_f_avg, gn_f_avg]):
        cell = ws.cell(row=rr, column=col, value=(None if np.isnan(v) else v))
        cell.number_format = "0.00"
        cell.alignment = RIGHT
    ws.cell(row=rr, column=12).fill = GNN_FILL

    def per_network_mean_fillgap(model_key: str) -> float:
        diffs = []
        for n in networks:
            b = lookup(agg, MODEL_LABELS[model_key], n, scen_short, "Mean_Fill_Rate")
            g = lookup(agg, MODEL_LABELS["gnn"],     n, scen_short, "Mean_Fill_Rate")
            if np.isnan(b) or np.isnan(g):
                continue
            diffs.append(g - b)
        return float(np.mean(diffs)) if diffs else float("nan")

    for col, mk in zip([13, 14, 15], ["basestock", "mappo", "happo"]):
        v = per_network_mean_fillgap(mk)
        cell = ws.cell(row=rr, column=col, value=(None if np.isnan(v) else v))
        cell.number_format = "0.00"
        cell.alignment = RIGHT
        if not np.isnan(v):
            cell.font = POS_FONT if v >= 0 else NEG_FONT

    for c in range(1, len(headers) + 1):
        ws.cell(row=rr, column=c).border = BORDER

    widths = [10] + [18] * 4 + [12] * 3 + [14] * 4 + [16] * 3
    autosize(ws, widths)
    ws.freeze_panes = "B4"


def build_gap_summary(ws: Worksheet, agg: pd.DataFrame, networks: list[str],
                      scenarios: dict) -> None:
    ws.cell(row=1, column=1, value="Gap-% Summary vs GNN-HAPPO"
            ).font = Font(bold=True, size=14)

    headers = ["Scenario", "Baseline",
               "Mean Cost Gap % (across networks)",
               "Min Cost Gap %", "Max Cost Gap %",
               "Mean Fill-Rate Gap (pp)",
               "Networks where GNN wins"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=3, column=j, value=h)
    style_header(ws, 3, len(headers))

    r = 4
    for sc_val in scenarios.values():
        scen = sc_val["short"]
        for mk in BASELINES:
            cost_gaps = []
            fill_gaps = []
            wins = 0
            n_valid = 0
            for n in networks:
                b = lookup(agg, MODEL_LABELS[mk], n, scen, "Mean_Total_Cost")
                g = lookup(agg, MODEL_LABELS["gnn"], n, scen, "Mean_Total_Cost")
                bf = lookup(agg, MODEL_LABELS[mk], n, scen, "Mean_Fill_Rate")
                gf = lookup(agg, MODEL_LABELS["gnn"], n, scen, "Mean_Fill_Rate")
                if np.isnan(b) or np.isnan(g) or b == 0:
                    continue
                gap = (b - g) / b * 100.0
                cost_gaps.append(gap)
                if gap > 0:
                    wins += 1
                n_valid += 1
                if not (np.isnan(bf) or np.isnan(gf)):
                    fill_gaps.append(gf - bf)

            ws.cell(row=r, column=1, value=sc_val["label"]).alignment = LEFT
            ws.cell(row=r, column=2, value=MODEL_LABELS[mk]).alignment = LEFT
            if cost_gaps:
                ws.cell(row=r, column=3, value=float(np.mean(cost_gaps))).number_format = "0.00"
                ws.cell(row=r, column=4, value=float(np.min(cost_gaps))).number_format = "0.00"
                ws.cell(row=r, column=5, value=float(np.max(cost_gaps))).number_format = "0.00"
                cell = ws.cell(row=r, column=3)
                cell.font = POS_FONT if cell.value >= 0 else NEG_FONT
            else:
                for col in (3, 4, 5):
                    ws.cell(row=r, column=col).fill = MISS_FILL
            if fill_gaps:
                fg_cell = ws.cell(row=r, column=6, value=float(np.mean(fill_gaps)))
                fg_cell.number_format = "0.00"
                fg_cell.font = POS_FONT if fg_cell.value >= 0 else NEG_FONT
            else:
                ws.cell(row=r, column=6).fill = MISS_FILL
            ws.cell(row=r, column=7, value=f"{wins}/{n_valid}").alignment = CENTER
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).border = BORDER
                if c >= 3 and c <= 6:
                    ws.cell(row=r, column=c).alignment = RIGHT
            r += 1

    autosize(ws, [30, 18, 22, 14, 14, 22, 22])
    ws.freeze_panes = "C4"


def build_raw(ws: Worksheet, df: pd.DataFrame) -> None:
    cols = ["Model", "Network", "Scenario", "Episode", "Status", "Note",
            "Total_Cost", "Fill_Rate", "Lost_Sales", "Avg_Inventory",
            "Total_Holding_Cost", "Total_Backlog_Cost", "Total_Ordering_Cost"]
    for j, h in enumerate(cols, start=1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, 1, len(cols))

    for i, rec in enumerate(df[cols].itertuples(index=False, name=None), start=2):
        for j, val in enumerate(rec, start=1):
            cell = ws.cell(row=i, column=j, value=(None if (isinstance(val, float) and np.isnan(val)) else val))
            if isinstance(val, float) and not np.isnan(val):
                cell.number_format = "#,##0.0000"
                cell.alignment = RIGHT
        if rec[4] != "OK":
            for c in range(1, len(cols) + 1):
                ws.cell(row=i, column=c).fill = MISS_FILL
    autosize(ws, [18, 8, 14, 8, 10, 30, 16, 12, 12, 14, 18, 18, 18])
    ws.freeze_panes = "E2"


def build_methodology(ws: Worksheet, num_episodes: int, episode_length: int,
                      seed: int) -> None:
    ws.cell(row=1, column=1, value="Methodology & Formulas"
            ).font = Font(bold=True, size=14)
    lines = [
        ("", ""),
        ("Models",  ", ".join(MODEL_LABELS[m] for m in MODEL_FAMILIES) + "   (GNN-HAPPO = proposed)"),
        ("Networks", ", ".join(NETWORKS)),
        ("Scenarios", ", ".join(v["short"] for v in DEMAND_SCENARIOS.values())),
        ("Episodes per cell", str(num_episodes)),
        ("Episode length (days)", str(episode_length)),
        ("Random seed", str(seed)),
        ("", ""),
        ("Per-episode aggregation",
            "For each (model, network, scenario), all 5 episode rows are averaged "
            "to obtain Mean Total Cost and Mean Fill Rate."),
        ("Total Cost formula",
            "Total_Cost = sum(Total_Holding_Cost + Total_Backlog_Cost + Total_Ordering_Cost) "
            "over all agents in the episode."),
        ("Fill Rate formula",
            "Mean over all agents of `service_level`. For retailers, `service_level = "
            "orders_from_stock / orders_placed x 100`. For DCs, the env-reported "
            "DC cycle service level is used."),
        ("Cost Gap %",
            "(Baseline_Cost - GNN_Cost) / Baseline_Cost x 100. Positive = GNN-HAPPO is cheaper."),
        ("Fill-Rate Gap (pp)",
            "GNN_FillRate - Baseline_FillRate. Positive = GNN-HAPPO serves more demand."),
        ("Average across networks (Gap_Summary, Scenario sheets)",
            "Per-network relative gaps are averaged with equal weight per network "
            "(small and large networks count the same)."),
        ("Demand scenario injection",
            "After the evaluator constructs its env, demand_mean[sku] and demand_std[sku] "
            "are overwritten with the scenario values. All other env parameters "
            "(costs, lead times, dc_assignments) come from the per-network YAML config."),
        ("Highlighting", "Green = GNN-HAPPO column / positive gap. Red = negative gap or MISSING. Yellow = best model in row."),
        ("Missing-model handling", "RL cells where results/<algo>_<net>/run_seed_1/models is absent are marked MISSING."),
    ]
    for i, (k, v) in enumerate(lines, start=2):
        ws.cell(row=i, column=1, value=k).font = Font(bold=bool(k))
        ws.cell(row=i, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=i, column=2, value=v).alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 110


def write_excel(df: pd.DataFrame, out_path: Path, networks: list[str],
                scenarios: dict, num_episodes: int, episode_length: int,
                seed: int) -> None:
    agg = aggregate(df)

    wb = Workbook()
    wb.remove(wb.active)
    build_overview(wb.create_sheet("Overview"), agg, networks, scenarios)
    for sc_val in scenarios.values():
        ws = wb.create_sheet(sc_val["sheet"])
        build_scenario_sheet(ws, agg, networks, sc_val["short"], sc_val["label"])
    build_gap_summary(wb.create_sheet("Gap_Summary"), agg, networks, scenarios)
    build_raw(wb.create_sheet("Raw_Data"), df)
    build_methodology(wb.create_sheet("Methodology"), num_episodes,
                      episode_length, seed)

    if out_path.exists():
        try:
            out_path.unlink()
        except PermissionError as e:
            raise PermissionError(
                f"Cannot delete {out_path}. Close the file in Excel and re-run."
            ) from e
    wb.save(out_path)
    print(f"\n[OK] Wrote {out_path}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def parse_cli():
    p = argparse.ArgumentParser(description="Network x Scenario benchmark")
    p.add_argument("--num_episodes",   type=int, default=5)
    p.add_argument("--episode_length", type=int, default=90)
    p.add_argument("--seed",           type=int, default=42)
    p.add_argument("--cuda",           action="store_true", default=False)
    p.add_argument("--models", nargs="+", default=MODEL_FAMILIES,
                   choices=MODEL_FAMILIES,
                   help="Subset of models to run.")
    p.add_argument("--networks", nargs="+", default=NETWORKS,
                   choices=NETWORKS,
                   help="Subset of networks to run.")
    p.add_argument("--scenarios", nargs="+",
                   default=list(DEMAND_SCENARIOS.keys()),
                   choices=list(DEMAND_SCENARIOS.keys()),
                   help="Subset of scenarios to run.")
    p.add_argument("--out", type=str,
                   default=str(ROOT / "Network_Scenario_Benchmark.xlsx"),
                   help="Output Excel path.")
    p.add_argument("--raw_csv", type=str, default=None,
                   help="Optional path to also save the raw per-episode CSV.")
    return p.parse_args()


def main():
    args = parse_cli()

    scenarios = {k: DEMAND_SCENARIOS[k] for k in args.scenarios}

    print("=" * 70)
    print("NETWORK x SCENARIO BENCHMARK")
    print("=" * 70)
    print(f"  Models    : {', '.join(MODEL_LABELS[m] for m in args.models)}")
    print(f"  Networks  : {', '.join(args.networks)}")
    print(f"  Scenarios : {', '.join(s for s in args.scenarios)}")
    print(f"  Episodes  : {args.num_episodes} per cell")
    print(f"  Ep length : {args.episode_length} days")
    print(f"  Seed      : {args.seed}")
    print(f"  CUDA      : {args.cuda}")
    print("=" * 70)

    df = run_all(args.models, args.networks, scenarios,
                 args.episode_length, args.num_episodes, args.seed, args.cuda)

    if args.raw_csv:
        Path(args.raw_csv).parent.mkdir(parents=True, exist_ok=True)
        df.to_csv(args.raw_csv, index=False)
        print(f"[OK] Raw per-episode CSV saved -> {args.raw_csv}")

    write_excel(df, Path(args.out), args.networks, scenarios,
                args.num_episodes, args.episode_length, args.seed)


if __name__ == "__main__":
    main()
