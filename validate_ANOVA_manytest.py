#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
ANOVA Validation Across 10 Networks x 4 Models
================================================

For each of 10 networks, runs a one-way ANOVA comparing the 4 models on
Total_Cost and Fill_Rate.  If the omnibus ANOVA is significant, runs
Tukey HSD post-hoc to identify which model pairs differ.

Hypotheses (per network, per metric):
    H0: mean(BaseStock) = mean(MAPPO) = mean(HAPPO) = mean(GNN-HAPPO)
    H1: at least one mean differs

Data source: evaluation_results/<model>_<net>/results_*_<net>.csv
             (5 episodes per file).

Output: Validate_ANOVA_ManyTest.xlsx with sheets:
    - ANOVA_Cost       : per-network omnibus ANOVA on Total_Cost
    - ANOVA_FillRate   : per-network omnibus ANOVA on Fill_Rate
    - Tukey_Cost       : pairwise post-hoc on Total_Cost (only sig networks)
    - Tukey_FillRate   : pairwise post-hoc on Fill_Rate (only sig networks)
    - Descriptive      : mean/std per (network, model) for both metrics
    - Summary          : compact win/loss table for GNN-HAPPO vs each baseline
    - Methodology      : formulas, hypotheses, interpretation

Usage:
    python validate_ANOVA_manytest.py
    python validate_ANOVA_manytest.py --alpha 0.01
"""

from __future__ import annotations

import argparse
from itertools import combinations
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from scipy import stats
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


ROOT = Path(__file__).resolve().parent
EVAL = ROOT / "evaluation_results"
OUTPUT = ROOT / "Validate_ANOVA_ManyTest.xlsx"

NETWORKS = ["1x3", "1x7", "1x10", "2x15", "2x20", "2x30", "2x40",
            "4x15", "4x30", "4x40"]

MODELS = ["(s,S) Policy", "MAPPO", "HAPPO", "GNN-HAPPO"]
BASELINES = ["(s,S) Policy", "MAPPO", "HAPPO"]
PROPOSED = "GNN-HAPPO"


# ---------------------------------------------------------------------------
# CSV path resolution
# ---------------------------------------------------------------------------

def csv_path(model: str, net: str) -> Path:
    if model == "(s,S) Policy":
        return EVAL / f"basestock_{net}" / f"results_ss_heuristic_{net}.csv"
    if model == "MAPPO":
        stem = "results_mappo" if net == "2x15" else "results_standard_mappo"
        return EVAL / f"mappo_{net}" / f"{stem}_{net}.csv"
    if model == "HAPPO":
        return EVAL / f"happo_{net}" / f"results_standard_happo_{net}.csv"
    if model == "GNN-HAPPO":
        return EVAL / f"gnn_{net}" / f"results_gnn_happo_{net}.csv"
    raise ValueError(model)


def load_episodes(model: str, net: str) -> pd.DataFrame:
    p = csv_path(model, net)
    if not p.exists():
        raise FileNotFoundError(p)
    return pd.read_csv(p)


# ---------------------------------------------------------------------------
# ANOVA + Tukey HSD
# ---------------------------------------------------------------------------

def run_anova_one(arrays: list[np.ndarray], group_names: list[str]) -> dict:
    """Run one-way ANOVA on k groups and compute SS table + effect size."""
    all_data = np.concatenate(arrays)
    grand_mean = float(np.mean(all_data))
    k = len(arrays)
    n_total = len(all_data)

    ss_between = float(sum(len(a) * (np.mean(a) - grand_mean) ** 2 for a in arrays))
    ss_within  = float(sum(np.sum((a - np.mean(a)) ** 2) for a in arrays))
    ss_total   = ss_between + ss_within

    df_between = k - 1
    df_within  = n_total - k
    df_total   = n_total - 1

    ms_between = ss_between / df_between if df_between > 0 else 0.0
    ms_within  = ss_within / df_within if df_within > 0 else 0.0

    F_stat, p_value = stats.f_oneway(*arrays)
    F_stat = float(F_stat)
    p_value = float(p_value)

    # Effect size: eta-squared = SS_between / SS_total
    eta_sq = ss_between / ss_total if ss_total > 0 else 0.0
    # Effect size interpretation (Cohen): 0.01 small, 0.06 medium, 0.14 large

    return {
        "k": k, "n_total": n_total,
        "ss_between": ss_between, "ss_within": ss_within, "ss_total": ss_total,
        "df_between": df_between, "df_within": df_within, "df_total": df_total,
        "ms_between": ms_between, "ms_within": ms_within,
        "F": F_stat, "p_value": p_value, "eta_sq": eta_sq,
    }


def run_tukey(arrays: list[np.ndarray], group_names: list[str], ms_within: float,
              df_within: int, alpha: float) -> list[dict]:
    """Run Tukey HSD post-hoc, return per-pair results."""
    try:
        res = stats.tukey_hsd(*arrays)
    except Exception as e:
        return [{"error": str(e)}]

    try:
        q_crit = stats.studentized_range.ppf(1 - alpha, len(arrays), df_within)
    except Exception:
        q_crit = None

    out = []
    for i, j in combinations(range(len(arrays)), 2):
        mean_i = float(np.mean(arrays[i]))
        mean_j = float(np.mean(arrays[j]))
        diff   = mean_i - mean_j
        pval   = float(res.pvalue[i, j])
        ci_low  = float(res.confidence_interval(confidence_level=1 - alpha).low[i, j])
        ci_high = float(res.confidence_interval(confidence_level=1 - alpha).high[i, j])
        if q_crit is not None and ms_within > 0:
            hsd_thresh = q_crit * np.sqrt(
                ms_within / 2.0 * (1.0 / len(arrays[i]) + 1.0 / len(arrays[j]))
            )
        else:
            hsd_thresh = float("nan")

        sig = pval < alpha
        out.append({
            "pair": f"{group_names[i]} vs {group_names[j]}",
            "model_A": group_names[i],
            "model_B": group_names[j],
            "mean_A": mean_i,
            "mean_B": mean_j,
            "mean_diff": diff,
            "abs_diff": abs(diff),
            "hsd_threshold": hsd_thresh,
            "p_value": pval,
            "ci_low": ci_low,
            "ci_high": ci_high,
            "significant": sig,
        })
    return out


def analyze_metric(metric: str, alpha: float) -> dict:
    """Run ANOVA + Tukey for each of 10 networks on one metric.

    Returns:
        {
            "per_network": [
                {"network": str, "anova": {...},
                 "tukey": [{...}], "descriptive": [{...}], "arrays": ..., "names": ...},
                ...
            ]
        }
    """
    per_network = []
    for net in NETWORKS:
        arrays = []
        names = []
        for m in MODELS:
            df = load_episodes(m, net)
            vals = df[metric].dropna().to_numpy(dtype=float)
            arrays.append(vals)
            names.append(m)

        desc = []
        for m, a in zip(names, arrays):
            desc.append({
                "model": m,
                "n": len(a),
                "mean": float(np.mean(a)),
                "std": float(np.std(a, ddof=1)) if len(a) > 1 else 0.0,
                "variance": float(np.var(a, ddof=1)) if len(a) > 1 else 0.0,
                "min": float(np.min(a)),
                "max": float(np.max(a)),
            })

        anova = run_anova_one(arrays, names)
        tukey = run_tukey(arrays, names, anova["ms_within"],
                          anova["df_within"], alpha) if anova["p_value"] < alpha else []

        per_network.append({
            "network": net,
            "descriptive": desc,
            "anova": anova,
            "tukey": tukey,
        })
    return {"per_network": per_network}


# ===========================================================================
# Excel builder
# ===========================================================================

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GNN_FILL = PatternFill("solid", fgColor="E2EFDA")
SIG_FILL = PatternFill("solid", fgColor="C6E0B4")
NSIG_FILL = PatternFill("solid", fgColor="FFF2CC")
WIN_FILL = PatternFill("solid", fgColor="E2EFDA")
LOSE_FILL = PatternFill("solid", fgColor="F4CCCC")
POS_FONT = Font(color="006100", bold=True)
NEG_FONT = Font(color="9C0006", bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws: Worksheet, row: int, n_cols: int):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def autosize(ws: Worksheet, widths: Iterable[int]):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def stars_for_p(p: float, alpha: float) -> str:
    if p < 0.001: return "***"
    if p < 0.01:  return "**"
    if p < alpha: return "*"
    return "n.s."


def effect_size_label(eta_sq: float) -> str:
    if eta_sq >= 0.14: return "Large"
    if eta_sq >= 0.06: return "Medium"
    if eta_sq >= 0.01: return "Small"
    return "Negligible"


def build_anova_sheet(ws: Worksheet, results: dict, metric: str, alpha: float):
    direction = ("lower is better" if metric == "Total_Cost"
                 else "higher is better")
    title = (f"One-Way ANOVA per Network — {metric} ({direction}, "
             f"alpha={alpha})")
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
    ws.cell(row=2, column=1,
            value="H0: all 4 model means are equal.  H1: at least one differs."
            ).font = Font(italic=True, color="595959")

    headers = [
        "Network", "N total", "k", "SS Between", "SS Within", "SS Total",
        "df Between", "df Within", "MS Between", "MS Within",
        "F-statistic", "p-value", "Sig.", "eta_sq", "Effect Size",
        "Conclusion",
    ]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=4, column=j, value=h)
    style_header(ws, 4, len(headers))

    for i, rec in enumerate(results["per_network"]):
        r = 5 + i
        a = rec["anova"]
        net = rec["network"]
        is_sig = a["p_value"] < alpha
        stars = stars_for_p(a["p_value"], alpha)
        conclusion = ("Reject H0 — model means differ"
                      if is_sig else "Fail to reject H0")

        ws.cell(row=r, column=1, value=net).alignment = CENTER
        ws.cell(row=r, column=2, value=a["n_total"]).alignment = CENTER
        ws.cell(row=r, column=3, value=a["k"]).alignment = CENTER

        for col, key in zip([4, 5, 6], ["ss_between", "ss_within", "ss_total"]):
            cell = ws.cell(row=r, column=col, value=a[key])
            cell.number_format = "#,##0.00"
            cell.alignment = RIGHT
        for col, key in zip([7, 8], ["df_between", "df_within"]):
            ws.cell(row=r, column=col, value=a[key]).alignment = CENTER
        for col, key in zip([9, 10], ["ms_between", "ms_within"]):
            cell = ws.cell(row=r, column=col, value=a[key])
            cell.number_format = "#,##0.00"
            cell.alignment = RIGHT

        ws.cell(row=r, column=11, value=a["F"]).number_format = "0.0000"
        ws.cell(row=r, column=11).alignment = RIGHT
        ws.cell(row=r, column=12, value=a["p_value"]).number_format = "0.000000"
        ws.cell(row=r, column=12).alignment = RIGHT

        sig_cell = ws.cell(row=r, column=13, value=stars)
        sig_cell.alignment = CENTER
        sig_cell.font = POS_FONT if is_sig else Font()

        ws.cell(row=r, column=14, value=a["eta_sq"]).number_format = "0.0000"
        ws.cell(row=r, column=14).alignment = RIGHT
        ws.cell(row=r, column=15, value=effect_size_label(a["eta_sq"])
                ).alignment = CENTER

        conc_cell = ws.cell(row=r, column=16, value=conclusion)
        conc_cell.alignment = LEFT
        row_fill = SIG_FILL if is_sig else NSIG_FILL
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = BORDER
        ws.cell(row=r, column=13).fill = row_fill
        ws.cell(row=r, column=16).fill = row_fill

    autosize(ws, [10, 8, 6, 14, 14, 14, 10, 10, 14, 14, 14, 14, 8, 10, 12, 32])
    ws.freeze_panes = "B5"


def build_tukey_sheet(ws: Worksheet, results: dict, metric: str, alpha: float):
    direction = ("lower is better" if metric == "Total_Cost"
                 else "higher is better")
    title = (f"Tukey HSD Post-Hoc — {metric} ({direction}, alpha={alpha})")
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
    ws.cell(row=2, column=1,
            value="Pairwise comparison shown ONLY for networks where omnibus "
                  "ANOVA was significant.  Family-wise error controlled at alpha."
            ).font = Font(italic=True, color="595959")

    headers = [
        "Network", "Comparison Pair", "Mean A", "Mean B", "Mean Diff (A-B)",
        "|Diff|", "HSD threshold", "p-value", "95% CI Low", "95% CI High",
        "Sig.", "Better Model",
    ]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=4, column=j, value=h)
    style_header(ws, 4, len(headers))

    row = 5
    for rec in results["per_network"]:
        net = rec["network"]
        if not rec["tukey"]:
            continue
        block_start = row
        for t in rec["tukey"]:
            if "error" in t:
                ws.cell(row=row, column=1, value=net).alignment = CENTER
                ws.cell(row=row, column=2, value=f"Error: {t['error']}")
                row += 1
                continue

            ws.cell(row=row, column=1, value=net).alignment = CENTER
            ws.cell(row=row, column=2, value=t["pair"]).alignment = LEFT
            for col, key in zip([3, 4, 5, 6, 7], ["mean_A", "mean_B",
                                                    "mean_diff", "abs_diff",
                                                    "hsd_threshold"]):
                cell = ws.cell(row=row, column=col, value=t[key])
                cell.number_format = "#,##0.00"
                cell.alignment = RIGHT
            ws.cell(row=row, column=8, value=t["p_value"]).number_format = "0.000000"
            ws.cell(row=row, column=8).alignment = RIGHT
            ws.cell(row=row, column=9, value=t["ci_low"]).number_format = "#,##0.00"
            ws.cell(row=row, column=9).alignment = RIGHT
            ws.cell(row=row, column=10, value=t["ci_high"]).number_format = "#,##0.00"
            ws.cell(row=row, column=10).alignment = RIGHT

            stars = stars_for_p(t["p_value"], alpha)
            sig_cell = ws.cell(row=row, column=11,
                               value=stars if t["significant"] else "n.s.")
            sig_cell.alignment = CENTER
            if t["significant"]:
                sig_cell.font = POS_FONT

            # Better model
            if t["significant"]:
                if metric == "Total_Cost":
                    better = t["model_A"] if t["mean_A"] < t["mean_B"] else t["model_B"]
                else:
                    better = t["model_A"] if t["mean_A"] > t["mean_B"] else t["model_B"]
            else:
                better = "—"
            better_cell = ws.cell(row=row, column=12, value=better)
            better_cell.alignment = CENTER
            if better == PROPOSED:
                better_cell.font = POS_FONT
                better_cell.fill = GNN_FILL

            for c in range(1, len(headers) + 1):
                ws.cell(row=row, column=c).border = BORDER
            row += 1

        if row > block_start + 1:
            ws.merge_cells(start_row=block_start, start_column=1,
                           end_row=row - 1, end_column=1)
            ws.cell(row=block_start, column=1).alignment = CENTER

    autosize(ws, [10, 30, 14, 14, 16, 12, 14, 14, 14, 14, 8, 16])
    ws.freeze_panes = "C5"


def build_descriptive_sheet(ws: Worksheet, results_cost: dict, results_fill: dict):
    ws.cell(row=1, column=1, value="Descriptive Statistics per (Network, Model)"
            ).font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value="N = 5 episodes per cell. Computed from "
                                    "results_*_<net>.csv files."
            ).font = Font(italic=True, color="595959")

    headers = [
        "Network", "Model", "N",
        "Cost Mean", "Cost Std", "Cost Variance", "Cost Min", "Cost Max",
        "Fill Mean", "Fill Std", "Fill Variance", "Fill Min", "Fill Max",
    ]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=4, column=j, value=h)
    style_header(ws, 4, len(headers))

    row = 5
    for rec_c, rec_f in zip(results_cost["per_network"],
                             results_fill["per_network"]):
        net = rec_c["network"]
        block_start = row
        for d_c, d_f in zip(rec_c["descriptive"], rec_f["descriptive"]):
            ws.cell(row=row, column=1, value=net).alignment = CENTER
            ws.cell(row=row, column=2, value=d_c["model"]).alignment = LEFT
            ws.cell(row=row, column=3, value=d_c["n"]).alignment = CENTER
            for col, key in zip([4, 5, 6, 7, 8],
                                ["mean", "std", "variance", "min", "max"]):
                cell = ws.cell(row=row, column=col, value=d_c[key])
                cell.number_format = "#,##0.00"
                cell.alignment = RIGHT
            for col, key in zip([9, 10, 11, 12, 13],
                                ["mean", "std", "variance", "min", "max"]):
                cell = ws.cell(row=row, column=col, value=d_f[key])
                cell.number_format = "0.0000"
                cell.alignment = RIGHT

            if d_c["model"] == PROPOSED:
                for c in range(1, len(headers) + 1):
                    ws.cell(row=row, column=c).fill = GNN_FILL
            for c in range(1, len(headers) + 1):
                ws.cell(row=row, column=c).border = BORDER
            row += 1

        ws.merge_cells(start_row=block_start, start_column=1,
                       end_row=row - 1, end_column=1)
        ws.cell(row=block_start, column=1).alignment = CENTER

    autosize(ws, [10, 16, 6, 14, 12, 14, 12, 12, 12, 10, 12, 10, 10])
    ws.freeze_panes = "C5"


def build_summary_sheet(ws: Worksheet, results_cost: dict, results_fill: dict,
                        alpha: float):
    ws.cell(row=1, column=1,
            value=f"Summary: GNN-HAPPO vs Baselines (Tukey HSD, alpha={alpha})"
            ).font = Font(bold=True, size=13)
    ws.cell(row=2, column=1,
            value="Each cell shows GNN-HAPPO's outcome vs the baseline in that "
                  "column for the given network. 'GNN wins' = Tukey HSD pair "
                  "is significant and GNN has the better mean."
            ).font = Font(italic=True, color="595959")

    def render_section(start_row: int, results: dict, metric: str,
                       direction_label: str) -> int:
        ws.cell(row=start_row, column=1,
                value=f"{metric.upper()} ({direction_label})"
                ).font = Font(bold=True, size=12)
        headers = ["Network", "ANOVA p-value", "Omnibus Sig.",
                   *[f"vs {bl}" for bl in BASELINES]]
        for j, h in enumerate(headers, start=1):
            ws.cell(row=start_row + 1, column=j, value=h)
        style_header(ws, start_row + 1, len(headers))

        wins_per_baseline = {bl: 0 for bl in BASELINES}
        for i, rec in enumerate(results["per_network"]):
            r = start_row + 2 + i
            net = rec["network"]
            anova_p = rec["anova"]["p_value"]
            omnibus_sig = anova_p < alpha
            stars = stars_for_p(anova_p, alpha)

            ws.cell(row=r, column=1, value=net).alignment = CENTER
            ws.cell(row=r, column=2, value=anova_p).number_format = "0.000000"
            ws.cell(row=r, column=2).alignment = RIGHT
            sig_cell = ws.cell(row=r, column=3, value=stars)
            sig_cell.alignment = CENTER
            if omnibus_sig:
                sig_cell.font = POS_FONT
                sig_cell.fill = SIG_FILL
            else:
                sig_cell.fill = NSIG_FILL

            # For each baseline column, check the GNN vs baseline Tukey pair
            for j, bl in enumerate(BASELINES, start=4):
                if not rec["tukey"]:
                    label = "ANOVA n.s."
                    fill = NSIG_FILL
                    font = Font()
                else:
                    # Find the GNN vs baseline pair (order may vary)
                    pair = None
                    for t in rec["tukey"]:
                        if "error" in t:
                            continue
                        if {t["model_A"], t["model_B"]} == {PROPOSED, bl}:
                            pair = t
                            break
                    if pair is None:
                        label = "n/a"
                        fill = NSIG_FILL
                        font = Font()
                    else:
                        if metric == "Total_Cost":
                            gnn_mean = (pair["mean_A"] if pair["model_A"] == PROPOSED
                                        else pair["mean_B"])
                            bl_mean = (pair["mean_B"] if pair["model_A"] == PROPOSED
                                       else pair["mean_A"])
                            diff_pct = ((gnn_mean - bl_mean) / bl_mean * 100.0
                                        if bl_mean != 0 else 0)
                            gnn_better = gnn_mean < bl_mean
                            diff_label = f"({diff_pct:+.1f}%)"
                        else:
                            gnn_mean = (pair["mean_A"] if pair["model_A"] == PROPOSED
                                        else pair["mean_B"])
                            bl_mean = (pair["mean_B"] if pair["model_A"] == PROPOSED
                                       else pair["mean_A"])
                            diff_pp = gnn_mean - bl_mean
                            gnn_better = gnn_mean > bl_mean
                            diff_label = f"({diff_pp:+.2f}pp)"
                        stars_p = stars_for_p(pair["p_value"], alpha)
                        if pair["significant"] and gnn_better:
                            label = f"GNN wins {diff_label} {stars_p}"
                            fill = WIN_FILL
                            font = POS_FONT
                            wins_per_baseline[bl] += 1
                        elif pair["significant"] and not gnn_better:
                            label = f"GNN loses {diff_label} {stars_p}"
                            fill = LOSE_FILL
                            font = NEG_FONT
                        else:
                            label = f"n.s. {diff_label}"
                            fill = NSIG_FILL
                            font = Font()

                cell = ws.cell(row=r, column=j, value=label)
                cell.alignment = CENTER
                cell.fill = fill
                cell.font = font
                cell.border = BORDER

            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).border = BORDER

        # Win count row
        r_count = start_row + 2 + len(results["per_network"])
        ws.cell(row=r_count, column=1,
                value="GNN wins / total").font = Font(bold=True)
        ws.cell(row=r_count, column=1).border = BORDER
        ws.cell(row=r_count, column=2, value="").border = BORDER
        ws.cell(row=r_count, column=3, value="").border = BORDER
        for j, bl in enumerate(BASELINES, start=4):
            cell = ws.cell(row=r_count, column=j,
                           value=f"{wins_per_baseline[bl]}/{len(results['per_network'])}")
            cell.alignment = CENTER
            cell.font = Font(bold=True)
            cell.border = BORDER

        return r_count + 3

    next_row = render_section(4, results_cost, "Total_Cost", "lower is better")
    render_section(next_row, results_fill, "Fill_Rate", "higher is better")

    autosize(ws, [10, 14, 12, 26, 26, 26])
    ws.freeze_panes = "B5"


def build_methodology(ws: Worksheet, alpha: float):
    ws.cell(row=1, column=1, value="Methodology").font = Font(bold=True, size=14)
    lines = [
        ("", ""),
        ("Test", "One-way ANOVA (scipy.stats.f_oneway) per network, "
                 "comparing 4 model groups."),
        ("Groups (k)", "k = 4 models: (s,S) Policy, MAPPO, HAPPO, GNN-HAPPO."),
        ("Sample size", "n = 5 episodes per model per network (same seed=42)."),
        ("", ""),
        ("Hypotheses (per network, per metric)",
            "H0: mu_BS = mu_MAPPO = mu_HAPPO = mu_GNN.  "
            "H1: at least one mean differs."),
        ("", ""),
        ("ANOVA decomposition",
            "SS_total = SS_between + SS_within.  "
            "F = MS_between / MS_within = (SS_between/(k-1)) / (SS_within/(N-k))."),
        ("Effect size", "eta_sq = SS_between / SS_total.  "
                          "Cohen interpretation: 0.01 small, 0.06 medium, 0.14 large."),
        ("", ""),
        ("Post-hoc", "Tukey HSD (scipy.stats.tukey_hsd) — controls family-wise "
                      "error rate across all pairwise comparisons.  Only run "
                      "when omnibus ANOVA is significant."),
        ("HSD threshold",
            "HSD = q_crit(alpha, k, df_within) * sqrt(MS_within/2 * (1/n_A + 1/n_B)). "
            "If |mean_A - mean_B| > HSD, the pair differs significantly."),
        ("", ""),
        ("Significance level", f"alpha = {alpha}"),
        ("Stars", "* p < alpha,  ** p < 0.01,  *** p < 0.001"),
        ("", ""),
        ("Interpretation guide", ""),
        ("  Green (GNN wins)",
            "Tukey HSD pair (GNN vs baseline) is significant AND GNN has the "
            "better mean (lower cost / higher fill rate)."),
        ("  Red (GNN loses)",
            "Tukey HSD pair is significant but baseline has the better mean."),
        ("  Yellow (n.s. / ANOVA n.s.)",
            "Either the omnibus ANOVA was not significant (no Tukey run) or the "
            "specific pair was not significant after HSD correction."),
        ("", ""),
        ("Sheet guide", ""),
        ("  ANOVA_Cost / ANOVA_FillRate",
            "Omnibus ANOVA: F-statistic, p-value, SS/df/MS table, eta-squared."),
        ("  Tukey_Cost / Tukey_FillRate",
            "Pairwise post-hoc comparisons (6 pairs per network) shown ONLY for "
            "networks where the omnibus was significant."),
        ("  Descriptive", "Mean/Std/Variance/Min/Max per (network, model) "
                           "for both metrics."),
        ("  Summary", "Compact win/lose/n.s. table for GNN-HAPPO vs each "
                       "baseline at every network."),
        ("", ""),
        ("Limitation",
            "n=5 per group → df_within = 16 for k=4.  Tukey HSD has limited "
            "power at this sample size; effect sizes must be large to reach "
            "significance.  Increasing to 30-100 episodes per model would "
            "strengthen conclusions."),
        ("Data source",
            "evaluation_results/<model>_<net>/results_*_<net>.csv "
            "(5 episodes, same seed=42)."),
    ]
    for i, (k, v) in enumerate(lines, start=2):
        ws.cell(row=i, column=1, value=k).font = Font(bold=bool(k))
        ws.cell(row=i, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=i, column=2, value=v).alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110


def write_excel(results_cost: dict, results_fill: dict, alpha: float):
    wb = Workbook()
    wb.remove(wb.active)
    build_anova_sheet(wb.create_sheet("ANOVA_Cost"), results_cost,
                      "Total_Cost", alpha)
    build_anova_sheet(wb.create_sheet("ANOVA_FillRate"), results_fill,
                      "Fill_Rate", alpha)
    build_tukey_sheet(wb.create_sheet("Tukey_Cost"), results_cost,
                      "Total_Cost", alpha)
    build_tukey_sheet(wb.create_sheet("Tukey_FillRate"), results_fill,
                      "Fill_Rate", alpha)
    build_descriptive_sheet(wb.create_sheet("Descriptive"),
                             results_cost, results_fill)
    build_summary_sheet(wb.create_sheet("Summary"), results_cost,
                         results_fill, alpha)
    build_methodology(wb.create_sheet("Methodology"), alpha)

    if OUTPUT.exists():
        try:
            OUTPUT.unlink()
        except PermissionError as e:
            raise PermissionError(
                f"Cannot delete {OUTPUT}. Close the file in Excel and re-run."
            ) from e
    wb.save(OUTPUT)
    print(f"\n[OK] Wrote {OUTPUT}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="ANOVA validation across 10 networks x 4 models")
    parser.add_argument("--alpha", type=float, default=0.05,
                        help="Significance level (default: 0.05)")
    args = parser.parse_args()

    print("=" * 60)
    print("ANOVA VALIDATION (10 networks x 4 models)")
    print("=" * 60)
    print(f"  Models   : {', '.join(MODELS)}")
    print(f"  Networks : {', '.join(NETWORKS)}")
    print(f"  alpha    : {args.alpha}")
    print("=" * 60)

    print("\nRunning ANOVA on Total_Cost...")
    results_cost = analyze_metric("Total_Cost", args.alpha)
    print("Running ANOVA on Fill_Rate...")
    results_fill = analyze_metric("Fill_Rate", args.alpha)

    # Console summary
    print(f"\n{'Network':<8} {'Cost F':>10} {'Cost p':>12} {'Sig':>5}  "
          f"{'Fill F':>10} {'Fill p':>12} {'Sig':>5}")
    print("-" * 70)
    for rc, rf in zip(results_cost["per_network"], results_fill["per_network"]):
        net = rc["network"]
        ac, af = rc["anova"], rf["anova"]
        sc = stars_for_p(ac["p_value"], args.alpha)
        sf = stars_for_p(af["p_value"], args.alpha)
        print(f"{net:<8} {ac['F']:>10.3f} {ac['p_value']:>12.6f} {sc:>5}  "
              f"{af['F']:>10.3f} {af['p_value']:>12.6f} {sf:>5}")

    write_excel(results_cost, results_fill, args.alpha)


if __name__ == "__main__":
    main()
