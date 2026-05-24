"""Build Model_Comparison_Across_Networks.xlsx from per-(model, network) episode CSVs.

Aggregates Total_Cost and Fill_Rate across episodes (mean) for each
(model, network) pair, then writes three sheets:
  - Raw Results
  - Average Performance
  - Gap Analysis vs GNN-HAPPO
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
EVAL = ROOT / "evaluation_results"

# (instance_label, network_size) in display order
INSTANCES = [
    ("Instance 1", "1x7"),
    ("Instance 2", "2x15 (base)"),
    ("Instance 3", "2x30"),
    ("Instance 4", "4x40"),
]

# model -> {network_key -> csv path}.  network_key matches the second element of INSTANCES.
CSV_PATHS = {
    "(s,S) Policy": {
        "1x7":         EVAL / "basestock_1x7"  / "results_ss_heuristic_1x7.csv",
        "2x15 (base)": EVAL / "basestock_2x15" / "results_ss_heuristic_2x15.csv",
        "2x30":        EVAL / "basestock_2x30" / "results_ss_heuristic_2x30.csv",
        "4x40":        EVAL / "basestock_4x40" / "results_ss_heuristic_4x40.csv",
    },
    "MAPPO": {
        "1x7":         EVAL / "mappo_1x7"  / "results_standard_mappo_1x7.csv",
        "2x15 (base)": EVAL / "mappo_2x15" / "results_mappo_2x15.csv",
        "2x30":        EVAL / "mappo_2x30" / "results_standard_mappo_2x30.csv",
        "4x40":        EVAL / "mappo_4x40" / "results_standard_mappo_4x40.csv",
    },
    "HAPPO": {
        "1x7":         EVAL / "happo_1x7"  / "results_standard_happo_1x7.csv",
        "2x15 (base)": EVAL / "happo_2x15" / "results_standard_happo_2x15.csv",
        "2x30":        EVAL / "happo_2x30" / "results_standard_happo_2x30.csv",
        "4x40":        EVAL / "happo_4x40" / "results_standard_happo_4x40.csv",
    },
    "GNN-HAPPO": {
        "1x7":         EVAL / "gnn_1x7"  / "results_gnn_happo_1x7.csv",
        "2x15 (base)": EVAL / "gnn_2x15" / "results_gnn_happo_2x15.csv",
        "2x30":        EVAL / "gnn_2x30" / "results_gnn_happo_2x30.csv",
        "4x40":        EVAL / "gnn_4x40" / "results_gnn_happo_4x40.csv",
    },
}

MODELS = ["(s,S) Policy", "MAPPO", "HAPPO", "GNN-HAPPO"]
OUTPUT = ROOT / "Model_Comparison_Across_Networks.xlsx"


def load_means() -> dict:
    """Return {model: {network: {'cost': mean_total_cost, 'fill': mean_fill_rate}}}."""
    out: dict = {}
    for model, by_net in CSV_PATHS.items():
        out[model] = {}
        for net, path in by_net.items():
            if not path.exists():
                raise FileNotFoundError(f"Missing CSV for {model} / {net}: {path}")
            df = pd.read_csv(path)
            out[model][net] = {
                "cost": float(df["Total_Cost"].mean()),
                "fill": float(df["Fill_Rate"].mean()),
            }
    return out


# ---------- styling helpers ----------

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
GNN_FILL = PatternFill("solid", fgColor="E2EFDA")  # highlight reference column
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header_row(ws, row: int, n_cols: int) -> None:
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def autosize(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------- sheet builders ----------

def build_raw(ws, means: dict) -> None:
    headers = ["Instance", "Network Size", "Metric", *MODELS]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    row = 2
    for inst_label, net in INSTANCES:
        # Total Cost row (display in VND thousands)
        ws.cell(row=row, column=1, value=inst_label)
        ws.cell(row=row, column=2, value=net)
        ws.cell(row=row, column=3, value="Total Cost (VND thousands)")
        for j, m in enumerate(MODELS, start=4):
            cell = ws.cell(row=row, column=j, value=means[m][net]["cost"] / 1000.0)
            cell.number_format = "#,##0.00"
        # Fill Rate row
        ws.cell(row=row + 1, column=1, value=inst_label)
        ws.cell(row=row + 1, column=2, value=net)
        ws.cell(row=row + 1, column=3, value="Fill Rate (%)")
        for j, m in enumerate(MODELS, start=4):
            cell = ws.cell(row=row + 1, column=j, value=means[m][net]["fill"])
            cell.number_format = "0.00"

        # styling for the two-row instance block
        for r in (row, row + 1):
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).border = BORDER
            ws.cell(row=r, column=1).alignment = LEFT
            ws.cell(row=r, column=2).alignment = CENTER
            ws.cell(row=r, column=3).alignment = LEFT
            for c in range(4, len(headers) + 1):
                ws.cell(row=r, column=c).alignment = RIGHT
            # highlight GNN-HAPPO column (last)
            ws.cell(row=r, column=len(headers)).fill = GNN_FILL
        # merge the instance / network labels across the two metric rows
        ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=1)
        ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        ws.cell(row=row, column=1).alignment = CENTER
        ws.cell(row=row, column=2).alignment = CENTER

        row += 2

    autosize(ws, [12, 14, 28, 14, 14, 14, 16])
    ws.freeze_panes = "D2"


def build_average(ws, means: dict) -> None:
    headers = ["Metric", *MODELS]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    # Average Total Cost (VND thousands)
    ws.cell(row=2, column=1, value="Avg Total Cost (VND thousands)")
    # Average Fill Rate (%)
    ws.cell(row=3, column=1, value="Avg Fill Rate (%)")

    for j, m in enumerate(MODELS, start=2):
        costs = [means[m][net]["cost"] for _, net in INSTANCES]
        fills = [means[m][net]["fill"] for _, net in INSTANCES]
        c1 = ws.cell(row=2, column=j, value=sum(costs) / len(costs) / 1000.0)
        c1.number_format = "#,##0.00"
        c2 = ws.cell(row=3, column=j, value=sum(fills) / len(fills))
        c2.number_format = "0.00"

    for r in (2, 3):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = BORDER
        ws.cell(row=r, column=1).alignment = LEFT
        for c in range(2, len(headers) + 1):
            ws.cell(row=r, column=c).alignment = RIGHT
        ws.cell(row=r, column=len(headers)).fill = GNN_FILL

    autosize(ws, [34, 14, 14, 14, 16])
    ws.freeze_panes = "B2"


def build_gap(ws, means: dict) -> None:
    baselines = ["(s,S) Policy", "MAPPO", "HAPPO"]
    headers = [
        "Instance",
        "Network Size",
        "Baseline",
        "Baseline Cost (VND thousands)",
        "GNN-HAPPO Cost (VND thousands)",
        "Absolute Cost Gap (VND thousands)",
        "Relative Cost Reduction (%)",
        "Baseline Fill Rate (%)",
        "GNN-HAPPO Fill Rate (%)",
        "Absolute Fill Rate Gap (pp)",
    ]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    row = 2
    for inst_label, net in INSTANCES:
        gnn_cost = means["GNN-HAPPO"][net]["cost"]
        gnn_fill = means["GNN-HAPPO"][net]["fill"]
        block_start = row
        for b in baselines:
            b_cost = means[b][net]["cost"]
            b_fill = means[b][net]["fill"]
            abs_cost_gap = b_cost - gnn_cost
            rel_red = (b_cost - gnn_cost) / b_cost * 100.0 if b_cost else 0.0
            abs_fill_gap = gnn_fill - b_fill

            ws.cell(row=row, column=1, value=inst_label)
            ws.cell(row=row, column=2, value=net)
            ws.cell(row=row, column=3, value=b)
            ws.cell(row=row, column=4, value=b_cost / 1000.0).number_format = "#,##0.00"
            ws.cell(row=row, column=5, value=gnn_cost / 1000.0).number_format = "#,##0.00"
            ws.cell(row=row, column=6, value=abs_cost_gap / 1000.0).number_format = "#,##0.00;[Red]-#,##0.00"
            ws.cell(row=row, column=7, value=rel_red).number_format = "0.00"
            ws.cell(row=row, column=8, value=b_fill).number_format = "0.00"
            ws.cell(row=row, column=9, value=gnn_fill).number_format = "0.00"
            ws.cell(row=row, column=10, value=abs_fill_gap).number_format = "0.00;[Red]-0.00"

            for c in range(1, len(headers) + 1):
                ws.cell(row=row, column=c).border = BORDER
            ws.cell(row=row, column=1).alignment = CENTER
            ws.cell(row=row, column=2).alignment = CENTER
            ws.cell(row=row, column=3).alignment = LEFT
            for c in range(4, len(headers) + 1):
                ws.cell(row=row, column=c).alignment = RIGHT
            row += 1

        # merge Instance + Network across the 3 baseline rows
        ws.merge_cells(start_row=block_start, start_column=1, end_row=row - 1, end_column=1)
        ws.merge_cells(start_row=block_start, start_column=2, end_row=row - 1, end_column=2)

    autosize(ws, [12, 14, 16, 22, 22, 22, 18, 18, 18, 18])
    ws.freeze_panes = "D2"


def main() -> None:
    means = load_means()
    wb = Workbook()
    wb.remove(wb.active)
    build_raw(wb.create_sheet("Raw Results"), means)
    build_average(wb.create_sheet("Average Performance"), means)
    build_gap(wb.create_sheet("Gap Analysis vs GNN-HAPPO"), means)
    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
